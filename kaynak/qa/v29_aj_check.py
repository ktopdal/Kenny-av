"""Alyssa v0.3–v0.22 uçtan uca testleri (Playwright, sahte Open-Meteo ve sahte GitHub API ile).
v0.29 genel/özel yayın ve Excel ile toplu envanter testleri
"""
import os, sys
import asyncio, base64, hashlib, json, os, re, shutil, subprocess, sys, time, urllib.parse
from playwright.async_api import async_playwright
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__))); import fbmock   # sahte Firebase

ROOT = '/home/claude/fishapp'
WWW = ROOT + '/qa/www_j'
SHOTS = ROOT + '/qa/v29j'
os.makedirs(SHOTS, exist_ok=True)
if os.path.isdir(WWW): shutil.rmtree(WWW)
shutil.copytree(os.environ.get('ALYSSA_DIST', ROOT + '/app/dist'), WWW + '/dist')   # ALYSSA_DIST: sabitlenmiş bir derlemeyi test etmek için
shutil.copytree(os.environ.get('ALYSSA_PUB', ROOT + '/app/dist_pub'), WWW + '/pub')   # genel yayın
prev = open(os.environ.get('ALYSSA_PREV', ROOT + '/app/preview/alyssa_onizleme.html'), encoding='utf-8').read()
open(WWW + '/preview.html', 'w', encoding='utf-8').write('<!doctype html><html><head><meta charset=utf-8><meta name=viewport content="width=device-width,initial-scale=1,viewport-fit=cover"></head><body>' + prev + '</body></html>')
PORT = 8799
srv = subprocess.Popen([sys.executable, '-m', 'http.server', str(PORT), '--bind', '127.0.0.1', '--directory', WWW], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
time.sleep(1)
BASE = 'http://127.0.0.1:%d' % PORT
KEY = json.load(open(ROOT + '/secret/key.json'))
m = re.search(r'<script id="sample" type="application/json">(.*?)</script>', prev, re.S)
SMP = json.loads(m.group(1))['raw']
FAKE_NOW = "const _D=Date;const _base=new _D('2026-10-17T09:30:00+03:00').getTime();const _s=_D.now();class FD extends _D{constructor(...a){if(a.length===0)super(_base+(_D.now()-_s));else super(...a);} static now(){return _base+(_D.now()-_s);}};window.Date=FD;"
calls = []
FAIL = []
def check(cond, msg):
    print(('OK   ' if cond else 'FAIL ') + msg)
    if not cond: FAIL.append(msg)

LAND = lambda lat, lon: 38.40 < lat < 38.50 and 27.15 < lon < 27.25   # sahte kara kutusu (test noktası 38.45, 27.20)
KIYI = lambda lat, lon: 38.60 < lat < 38.62 and 26.90 < lon < 26.92   # karada ama kıyı çizgisine 150 m (kıyıdan av)
LAND2 = lambda lat, lon: 38.36 < lat < 38.37 and 26.87 < lon < 26.88  # Kenny'nin ekran görüntüsü: karada, Overpass yanıt vermiyor, EMODnet 13 m diyor
TRAP = lambda lat, lon: 38.30 < lat < 38.32 and 26.70 < lon < 26.72   # denizde ama Nominatim yakın bir adres veriyor, derinlik servisi hata veriyor (eski sürüm "Kara" derdi)
HANG = {'gfs': False}
CORS = {'Access-Control-Allow-Origin': '*'}

async def mock_meteo(route):
    u = route.request.url; q = urllib.parse.parse_qs(urllib.parse.urlparse(u).query)
    calls.append(u)
    lat, lon = float(q['latitude'][0]), float(q['longitude'][0])
    model = (q.get('models') or [''])[0]
    if '/v1/elevation' in u:   # Copernicus 90 m yükselti: denizde 0
        el = 45 if LAND(lat, lon) else 38 if LAND2(lat, lon) else 4 if KIYI(lat, lon) else 0
        return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps({'elevation': [el]}))
    if 'marine' in u:
        hv = q.get('hourly', [''])[0].split(',')
        if LAND(lat, lon) or KIYI(lat, lon) or LAND2(lat, lon):
            if model:
                return await route.fulfill(status=400, content_type='application/json', headers=CORS, body=json.dumps({'error': True, 'reason': 'No data is available for this location'}))
            h = {'time': SMP['mbest']['hourly']['time']}
            for k in hv: h[k] = [None] * len(h['time'])
            return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps({'hourly': h, 'hourly_units': {}}))
        if model == 'meteofrance_wave':
            return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps(SMP['mf']))
        h = dict(SMP['mbest']['hourly']); h.update({k: SMP['mf']['hourly'][k] for k in ('wave_height', 'wave_direction', 'wave_period')})
        h = {k: v for k, v in h.items() if k == 'time' or k in hv}
        return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps({'hourly': h, 'hourly_units': SMP['mbest']['hourly_units']}))
    src = {'icon_eu': 'icon', 'ecmwf_ifs025': 'ecmwf', 'gfs_seamless': 'gfs', 'gfs_global': 'gfs', '': 'best'}.get(model)
    if src == 'gfs' and HANG['gfs']:
        return   # hiç yanıt verme: zaman aşımı testi
    if not src:
        return await route.fulfill(status=400, headers=CORS, body='{}')
    r = SMP[src]
    if src == 'gfs':   # örnek veride GFS yalnızca hamle: rüzgârı ECMWF'den %15 fazlasıyla türet (karşılaştırma satırları için)
        e = SMP['ecmwf']['hourly']; n = len(e['time']); gg = SMP['gfs']['hourly']['wind_gusts_10m']
        h = {k: (list(v) if k == 'time' else [None if x is None else (x * 1.15 if k == 'wind_speed_10m' else x) for x in v]) for k, v in e.items()}
        h['wind_gusts_10m'] = (gg + [None] * n)[:n]
        r = {'hourly': h, 'hourly_units': {'wind_speed_10m': 'm/s', 'wind_gusts_10m': (SMP['gfs'].get('hourly_units') or {}).get('wind_gusts_10m', 'm/s')}}
    if src == 'best':
        h = dict(SMP['ecmwf']['hourly']); h['precipitation_probability'] = SMP['best']['hourly']['precipitation_probability']; h['wind_gusts_10m'] = SMP['gfs']['hourly']['wind_gusts_10m']
        r = {'hourly': h, 'hourly_units': {'wind_speed_10m': 'm/s'}}
    return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps(r))

# ---- sahte GitHub Contents API (bellekte)
GH = {}   # path -> {'bytes': b, 'sha': s}
GH_LOG = []
GOOD_TOKEN = 'test-anahtar-gecerli'
async def mock_gh(route):
    req = route.request
    hdr = {'Access-Control-Allow-Origin': '*', 'Access-Control-Allow-Headers': 'authorization,accept,content-type,x-github-api-version', 'Access-Control-Allow-Methods': 'GET,PUT,OPTIONS'}
    if req.method == 'OPTIONS':
        return await route.fulfill(status=204, headers=hdr)
    p = urllib.parse.urlparse(req.url).path
    mm = re.match(r'/repos/([^/]+/[^/]+)/contents/(.+)$', p)
    GH_LOG.append((req.method, p))
    if not mm: return await route.fulfill(status=404, headers=hdr, body='{}')
    repo, path = mm.group(1), urllib.parse.unquote(mm.group(2))
    if req.headers.get('authorization') != 'Bearer ' + GOOD_TOKEN:
        return await route.fulfill(status=401, headers=hdr, content_type='application/json', body='{"message":"Bad credentials"}')
    if repo != 'kenny/alyssa-yedek':
        return await route.fulfill(status=404, headers=hdr, content_type='application/json', body='{"message":"Not Found"}')
    if req.method == 'GET':
        f = GH.get(path)
        if not f:
            kids = [k for k in GH if k.startswith(path + '/') and '/' not in k[len(path) + 1:]]
            if kids:
                return await route.fulfill(status=200, headers=hdr, content_type='application/json', body=json.dumps([{'name': k.split('/')[-1], 'path': k, 'sha': GH[k]['sha'], 'size': len(GH[k]['bytes']), 'type': 'file'} for k in sorted(kids)]))
            return await route.fulfill(status=404, headers=hdr, content_type='application/json', body='{"message":"Not Found"}')
        if 'raw' in (req.headers.get('accept') or ''):
            return await route.fulfill(status=200, headers=hdr, body=f['bytes'], content_type='application/octet-stream')
        return await route.fulfill(status=200, headers=hdr, content_type='application/json', body=json.dumps({'sha': f['sha'], 'path': path, 'size': len(f['bytes']), 'content': base64.b64encode(f['bytes']).decode(), 'encoding': 'base64'}))
    if req.method == 'PUT':
        body = json.loads(req.post_data)
        cur = GH.get(path)
        if cur and body.get('sha') != cur['sha']:
            return await route.fulfill(status=409, headers=hdr, content_type='application/json', body='{"message":"sha mismatch"}')
        if not cur and body.get('sha'):
            return await route.fulfill(status=422, headers=hdr, content_type='application/json', body='{"message":"sha given but file missing"}')
        b = base64.b64decode(body['content'])
        sha = hashlib.sha1(b).hexdigest()
        GH[path] = {'bytes': b, 'sha': sha}
        return await route.fulfill(status=201 if not cur else 200, headers=hdr, content_type='application/json', body=json.dumps({'content': {'sha': sha, 'path': path}}))
    return await route.fulfill(status=405, headers=hdr, body='{}')

PNG1 = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII=')
MAP_LOG = []
async def mock_tiles(route):
    MAP_LOG.append(route.request.url)
    return await route.fulfill(status=200, content_type='image/png', body=PNG1, headers=CORS)

async def mock_geo(route):
    u = route.request.url; MAP_LOG.append(u)
    if 'nominatim' in u:
        q = urllib.parse.parse_qs(urllib.parse.urlparse(u).query); lat, lon = float(q['lat'][0]), float(q['lon'][0])
        if LAND(lat, lon):
            return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps({'lat': str(lat), 'lon': str(lon), 'name': '', 'address': {'village': 'Test Köyü', 'province': 'İzmir'}}))
        if TRAP(lat, lon):
            return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps({'lat': str(lat + 0.001), 'lon': str(lon), 'name': '', 'address': {'suburb': 'Tuzak Mahallesi', 'province': 'İzmir'}}))
        if KIYI(lat, lon):
            return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps({'lat': str(lat), 'lon': str(lon), 'name': '', 'address': {'village': 'Kıyıköy', 'province': 'İzmir'}}))
        if LAND2(lat, lon):
            return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps({'lat': str(lat + 0.0003), 'lon': str(lon), 'name': '', 'address': {'suburb': 'Çelebi', 'province': 'İzmir'}}))
        return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps({'error': 'Unable to geocode'}))
    if 'overpass' in u:
        body = urllib.parse.unquote_plus(route.request.post_data or '')
        mm = re.search(r'around:6000,([-\d.]+),([-\d.]+)', body); lat, lon = float(mm.group(1)), float(mm.group(2))
        if TRAP(lat, lon): await asyncio.sleep(2.5)   # yavaş Overpass: ad ve diğer bilgiler önce gelmeli
        if LAND2(lat, lon): return await route.fulfill(status=504, headers=CORS, body='{}')   # Overpass yanıt vermedi
        cy = lat - 0.02 if LAND(lat, lon) else lat - 0.00135 if KIYI(lat, lon) else lat + 0.02   # OSM kıyı çizgisi batıdan doğuya: kara solda (kuzey), deniz sağda (güney)
        els = [{'type': 'way', 'id': 1, 'tags': {'natural': 'coastline'}, 'geometry': [{'lat': cy, 'lon': lon - 0.05}, {'lat': cy, 'lon': lon + 0.05}]},
               {'type': 'node', 'id': 2, 'lat': lat + 0.03, 'lon': lon, 'tags': {'place': 'village', 'name': 'Deneme'}},
               {'type': 'node', 'id': 3, 'lat': lat + 0.09, 'lon': lon, 'tags': {'place': 'town', 'name': 'Uzakkasaba'}}]
        return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps({'elements': els}))
    if 'emodnet-bathymetry.eu/depth_sample' in u:
        g = urllib.parse.unquote(urllib.parse.parse_qs(urllib.parse.urlparse(u).query)['geom'][0]); mm = re.search(r'POINT\(([-\d.]+) ([-\d.]+)\)', g)
        lon, lat = float(mm.group(1)), float(mm.group(2))
        if LAND(lat, lon): return await route.fulfill(status=404, headers=CORS, body='{}')
        if TRAP(lat, lon): return await route.fulfill(status=503, headers=CORS, body='{}')
        if KIYI(lat, lon): return await route.fulfill(status=404, headers=CORS, body='{}')
        if LAND2(lat, lon): return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps({'min': 12.1, 'max': 14.0, 'avg': 13.1, 'smoothed': 13.0}))   # karada yanlış derinlik
        return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps({'min': 53.1, 'max': 55.2, 'avg': 54.3, 'smoothed': 54.0}))
    return await route.fulfill(status=404, headers=CORS, body='{}')

# ---- sahte EUSeaMap (EMODnet Seabed Habitats WMS): GetFeatureInfo noktaya göre, diğerleri resim
HAB_LOG = []
HAB_FAIL = lambda lat, lon: 38.27 < lat < 38.28 and 26.62 < lon < 26.63   # servis yanıt vermiyor (CORS/bağlantı)
async def mock_hab(route):
    u = route.request.url; HAB_LOG.append(u)
    if 'GetFeatureInfo' not in u:
        return await route.fulfill(status=200, content_type='image/png', body=PNG1, headers=CORS)
    q = urllib.parse.parse_qs(urllib.parse.urlparse(u).query)
    x0, y0, x1, y1 = [float(v) for v in q['bbox'][0].split(',')]; lat, lon = (y0 + y1) / 2, (x0 + x1) / 2
    if HAB_FAIL(lat, lon): return await route.abort()
    feats = []
    if LAND(lat, lon) or LAND2(lat, lon) or KIYI(lat, lon): feats = []
    elif 38.24 < lat < 38.26 and 26.64 < lon < 26.66:
        feats = [{'type': 'Feature', 'properties': {'eunis2019c': 'MB252', 'eunis2019d': 'Posidonia oceanica meadows', 'biozone': 'Infralittoral', 'substrate': 'Seabed'}}]
    elif 38.21 < lat < 38.22 and 26.60 < lon < 26.62:
        feats = [{'type': 'Feature', 'properties': {'eunis2019c': 'MC15', 'eunis2019d': 'Circalittoral rock', 'biozone': 'Circalittoral', 'substrate': 'Rock or other hard substrata'}}]
    else:
        feats = [{'type': 'Feature', 'properties': {'eunis2019c': 'MB52', 'eunis2019d': 'Mediterranean infralittoral sand', 'biozone': 'Infralittoral', 'substrate': 'Sand', 'energy': 'Low energy'}}]
    return await route.fulfill(status=200, content_type='application/json', headers=CORS, body=json.dumps({'type': 'FeatureCollection', 'features': feats}))

# ---- sahte Claude API (Usta'ya sor)
CLAUDE_LOG = []
async def mock_claude(route):
    req = route.request
    hdr = {'Access-Control-Allow-Origin': '*', 'Access-Control-Allow-Headers': 'content-type,x-api-key,authorization,anthropic-version,anthropic-dangerous-direct-browser-access,anthropic-workspace-id', 'Access-Control-Allow-Methods': 'POST,OPTIONS'}
    if req.method == 'OPTIONS': return await route.fulfill(status=204, headers=hdr)
    body = json.loads(req.post_data or '{}'); CLAUDE_LOG.append({'h': dict(req.headers), 'b': body})
    key = req.headers.get('x-api-key')
    if key == 'sk-ant-test-nocredit':
        return await route.fulfill(status=400, headers=hdr, content_type='application/json', body=json.dumps({'type': 'error', 'error': {'type': 'invalid_request_error', 'message': 'Your credit balance is too low to access the Anthropic API. Please go to Plans & Billing to upgrade or purchase credits.'}}))
    if key == 'sk-ant-test-noeffort' and 'output_config' in body:
        return await route.fulfill(status=400, headers=hdr, content_type='application/json', body=json.dumps({'type': 'error', 'error': {'type': 'invalid_request_error', 'message': 'output_config: Extra inputs are not permitted'}}))
    if key not in ('sk-ant-test-good', 'sk-ant-test-noeffort'):
        return await route.fulfill(status=401, headers=hdr, content_type='application/json', body=json.dumps({'type': 'error', 'error': {'type': 'authentication_error', 'message': 'invalid x-api-key'}}))
    n = len([m for m in body.get('messages', []) if m['role'] == 'user'])
    txt = '**Önce en üstteki rigi bağla.**\n- Kurşun: kartta yazan kadar\n- Köstek: 0,40 mm\nGüvenlik: deniz D ise çıkma.' if n == 1 else 'İkinci cevap: akşam penceresini bekle.'
    return await route.fulfill(status=200, headers=hdr, content_type='application/json', body=json.dumps({'id': 'msg_test', 'type': 'message', 'role': 'assistant', 'model': body.get('model'),
        'content': [{'type': 'thinking', 'thinking': 'GİZLİ DÜŞÜNCE', 'signature': 'x'}, {'type': 'text', 'text': txt}], 'stop_reason': 'end_turn', 'usage': {'input_tokens': 3200, 'output_tokens': 420}}))

async def new_ctx(b, storage=None, w=390):
    ctx = await b.new_context(viewport={'width': w, 'height': 860}, timezone_id='Europe/Istanbul', device_scale_factor=2, permissions=['geolocation'], geolocation={'latitude': 38.30, 'longitude': 26.80})
    await ctx.add_init_script(FAKE_NOW)
    if storage:
        await ctx.add_init_script('if(!sessionStorage.getItem("seeded")){' + ''.join('localStorage.setItem(%s,%s);' % (json.dumps(k), json.dumps(v)) for k, v in storage.items()) + 'sessionStorage.setItem("seeded","1");}')
    await ctx.route('https://api.open-meteo.com/**', mock_meteo); await ctx.route('https://marine-api.open-meteo.com/**', mock_meteo)
    await ctx.route('https://api.github.com/**', mock_gh)
    await fbmock.install(ctx)
    for pat in ['https://tile.openstreetmap.org/**', 'https://server.arcgisonline.com/**', 'https://tiles.openseamap.org/**', 'https://ows.emodnet-bathymetry.eu/**', 'https://drive.emodnet-geology.eu/**']:
        await ctx.route(pat, mock_tiles)
    for pat in ['https://nominatim.openstreetmap.org/**', 'https://overpass-api.de/**', 'https://overpass.kumi.systems/**', 'https://rest.emodnet-bathymetry.eu/**']:
        await ctx.route(pat, mock_geo)
    await ctx.route('https://wa.me/**', lambda route: route.fulfill(status=200, content_type='text/html', body='<p>wa</p>'))
    await ctx.route('https://ows.emodnet-seabedhabitats.eu/**', mock_hab); await ctx.route('https://api.anthropic.com/**', mock_claude)
    pg = await ctx.new_page(); errs = []
    pg.on('pageerror', lambda e: errs.append('pageerror: ' + str(e)))
    pg.on('console', lambda m: errs.append('console: ' + m.text) if m.type == 'error' and 'fonts' not in m.text and 'ERR_' not in m.text and 'Failed to load resource' not in m.text else None)
    return ctx, pg, errs

def fake_now(iso):
    return FAKE_NOW.replace("2026-10-17T09:30:00+03:00", iso)

async def new_ctx_at(b, iso, w=390):
    ctx = await b.new_context(viewport={'width': w, 'height': 860}, timezone_id='Europe/Istanbul', device_scale_factor=2, permissions=['geolocation'], geolocation={'latitude': 38.30, 'longitude': 26.80})
    await ctx.add_init_script(fake_now(iso))
    await ctx.route('https://api.open-meteo.com/**', mock_meteo); await ctx.route('https://marine-api.open-meteo.com/**', mock_meteo)
    pg = await ctx.new_page(); errs = []
    pg.on('pageerror', lambda e: errs.append('pageerror: ' + str(e)))
    return ctx, pg, errs

async def overflow(pg, tag):
    ov = await pg.evaluate('document.documentElement.scrollWidth - window.innerWidth')
    check(ov <= 0, 'yatay taşma yok: ' + tag + (' (' + str(ov) + ' px)' if ov > 0 else ''))

async def row_idx(pg, name):
    return await pg.evaluate('(n) => [...document.querySelectorAll("#hTable tr")].findIndex(tr => { const b = tr.querySelector("th b"); return b && b.textContent === n && tr.querySelectorAll("td").length > 2; })', name)

def U(t): return t.upper().replace('İ', 'I')

async def state(pg):
    return await pg.evaluate('JSON.parse(localStorage.getItem("alyssa-v1"))')

# Kenny'nin kurulumu ve kalem adları özel veriden okunur; depoda düz yazılmaz
DB = json.load(open(ROOT + '/data/kenny_av_db.json', encoding='utf-8'))
_L1 = [l for l in DB['locations'] if l['id'] == '1'][0]
LDDM = ('%d°%s′ K' % (int(_L1['lat']), ('%06.3f' % ((_L1['lat'] % 1) * 60)).replace('.', ',')))   # plotter biçimi, özel dosyadan
SU = {s['id']: s for s in DB['setups']}; IT = {i['id']: i for i in DB['items']}
def imax(i): m = re.search(r'max\s*(\d+)\s*g', IT[i].get('spec', '')); return int(m.group(1)) if m else None
CAST = [i for i in IT if IT[i]['cat'] == 'rod' and 'tetikli' in IT[i].get('spec', '')]
PG_ROD, PG_REEL = SU['PG']['rods'][0], SU['PG']['reel']; TG_ROD, TG_REEL = SU['TG']['rods'][0], SU['TG']['reel']; NP_ROD = SU['NP']['rods'][0]
NM = lambda i: IT[i]['name']
OTHER_RODS = [i for i in IT if IT[i]['cat'] == 'rod' and i != PG_ROD]; OTHER_REELS = [i for i in IT if IT[i]['cat'] == 'reel' and i != PG_REEL]

async def tall(pg, h, path, sel=None):
    await pg.set_viewport_size({'width': 390, 'height': h}); await pg.wait_for_timeout(300)
    if sel: await pg.locator(sel).first.screenshot(path=path)
    else: await pg.screenshot(path=path)
    await pg.set_viewport_size({'width': 390, 'height': 860}); await pg.wait_for_timeout(150)

async def rigs_all(pg):
    await pg.click('.tabbar button[data-v="rigler"]'); await pg.wait_for_timeout(250)
    if await pg.locator('#sysBar button[data-sys="all"]').count(): await pg.click('#sysBar button[data-sys="all"]'); await pg.wait_for_timeout(250)

# ================================================== AJ) v0.29: genel / özel yayın, yeni ad, Excel ile toplu envanter
import openpyxl

async def lskeys(pg):
    return await pg.evaluate('Object.keys(localStorage)')

async def import_file(pg, path):
    await pg.click('.tabbar button[data-v="envanter"]'); await pg.wait_for_timeout(250)
    await pg.set_input_files('#xlFile', path); await pg.wait_for_timeout(900)
    return await pg.inner_text('#sheet')

async def main():
    async with async_playwright() as p:
        b = await p.chromium.launch(executable_path='/opt/pw-browsers/chromium-1194/chrome-linux/chrome', args=['--no-sandbox'])
        # ---------------- genel yayın
        src = open(WWW + '/pub/index.html', encoding='utf-8').read()
        bad = [w for w in ['Özel paket', 'özel paket', 'ozel.enc', 'Kenny', 'kenny', 'Seferihisar', 'Barbossa', 'ktopdal', 'Konum 1', "32'lik", '32 rig', 'Sığacık', 'paket şifre', 'Paket şifresi'] if w in src]
        check(not bad, 'AJ: genel yayında özel izler yok %s' % bad)
        check('// ----' not in src and '/*@' not in src, 'AJ: genel yayında kaynak yorumları ayıklandı')
        check(not os.path.exists(WWW + '/pub/paket'), 'AJ: genel yayında paket klasörü yok')
        man = json.load(open(WWW + '/pub/manifest.webmanifest', encoding='utf-8'))
        check(man['name'] == 'Alyssa WF' and man['short_name'] == 'Alyssa WF' and 'Weather & Fishing' in man.get('description', ''), 'AJ: manifest adı Alyssa WF')
        sw = open(WWW + '/pub/sw.js').read(); sw2 = open(WWW + '/dist/sw.js').read()
        check("P='awf-'" in sw and "P='alyssa-'" in sw2 and 'k.indexOf(P)===0' in sw, 'AJ: iki yayının önbellekleri ayrı (yalnız kendi önbelleğini siler)')
        rd = open(WWW + '/pub/README.md', encoding='utf-8').read()
        check('paket' not in rd.lower() and 'Alyssa WF' in rd, 'AJ: genel README paketsiz')
        ctx, pg, errs = await new_ctx(b)
        await pg.goto(BASE + '/pub/?qa=1&paket'); await pg.wait_for_timeout(1200)
        check(await pg.title() == 'Alyssa WF', 'AJ: başlık Alyssa WF')
        wel = await pg.inner_text('#welcome')
        check(await pg.locator('button[data-wel="paket"]').count() == 0 and '#welPkg' and await pg.locator('#welPkg').count() == 0, 'AJ: genel: karşılamada paket yok (?paket ile bile)')
        await pg.click('button[data-wel="bos"]'); await pg.wait_for_timeout(600)
        ks = await lskeys(pg)
        check('awf-v1' in ks and not any(k.startswith('alyssa-') or k.startswith('kenny') for k in ks), 'AJ: genel: depolama awf- önekiyle (%s)' % ks)
        await pg.click('.tabbar button[data-v="rigler"]'); await pg.wait_for_timeout(300)
        t = await pg.inner_text('#v-rigler')
        check('Özel paket' not in t and await pg.locator('#v-rigler [data-open="paket"]').count() == 0, 'AJ: genel: Rigler boş ekranında paket düğmesi yok')
        await pg.click('.tabbar button[data-v="daha"]'); await pg.wait_for_timeout(300)
        for i in range(6):
            await pg.click('#about'); await pg.wait_for_timeout(60)
        await pg.wait_for_timeout(300)
        check(await pg.locator('#moreMenu button[data-open="paket"]').count() == 0 and 'özel paket' not in (await pg.inner_text('#about')).lower(), 'AJ: genel: sürüm yazısına 6 dokunuş paket açmaz')
        check((await pg.inner_text('#about')).startswith('Alyssa WF v0.29'), 'AJ: sürüm yazısı Alyssa WF')
        await pg.click('#moreMenu button[data-open="ayarlar"]'); await pg.wait_for_timeout(300)
        sh = await pg.inner_text('#sheet')
        check('Özel paket' not in sh and 'Yeni kullanıcı gibi başla' not in sh and 'Profilimi sıfırla' in sh, 'AJ: genel: Sıfırla’da paket seçeneği yok')
        await pg.locator('#sheet [data-close]').first.click(); await pg.wait_for_timeout(200)
        # ---- örnek Excel
        await pg.click('.tabbar button[data-v="envanter"]'); await pg.wait_for_timeout(300)
        check(await pg.is_visible('#v-envanter button[data-xl="tpl"]') and await pg.is_visible('#v-envanter button[data-xl="pick"]'), 'AJ: Envanter’de Excel düğmeleri')
        async with pg.expect_download() as dl_info:
            await pg.click('#v-envanter button[data-xl="tpl"]')
        dl = await dl_info.value; tpl = SHOTS + '/ornek.xlsx'; await dl.save_as(tpl)
        check(dl.suggested_filename == 'ornek-envanter.xlsx', 'AJ: örnek dosya adı')
        wb = openpyxl.load_workbook(tpl)
        check(wb.sheetnames[:3] == ['Envanter', 'Nasıl doldurulur', 'Örnekler'] and wb['Listeler'].sheet_state == 'hidden', 'AJ: örnek dosyanın sayfaları (%s)' % wb.sheetnames)
        hd = [c.value for c in wb['Envanter'][1]]
        check(hd[:7] == ['Kategori', 'Marka', 'Model / ad', 'Özellik', 'Boylar', 'Durum', 'Not'] and 'Kamış tipi' in hd and 'Makine boyu' in hd, 'AJ: başlıklar yan yana')
        check(len(wb['Envanter'].data_validations.dataValidation) >= 5 and wb['Envanter'].max_row >= 2 and all(not c.value for c in wb['Envanter'][2]), 'AJ: açılır listeler var, Envanter sayfası boş')
        check(wb['Örnekler'].max_row >= 10 and wb['Örnekler']['E2'].number_format == '@', 'AJ: örnekler ve metin biçimli boy sütunu')
        # ---- doldur ve yükle
        ws = wb['Envanter']
        rows = [
            ['İğne', 'Owner', 'SSW 5111', 'siyah nikel', '1/0, 2/0, 3/0', 'Var', 'mercan için', 'Octopus / beak'],
            ['iğne', 'Gamakatsu', 'Chinu', '', 'No:2; No:4', 'Alınacak', '', 'chinu'],
            ['Florokarbon', 'Seaguar', 'Grand Max FX', '0,33 mm', '0,40 mm', ''],
            ['Kamış', 'Testrod', 'Slow 198', '', '', 'var', '', '', 'Slow jig', '1,98', 200, 2],
            ['Makine', 'Testreel', 'Lever 30', '', '', 'Alınacak', '', '', '', '', '', '', 'Çıkrık (lever drag / baitcasting)', 30],
            [None] * 6,
            ['Kurşun', '', 'Armut kurşun', '', '60 g, 80 g', 'Yok'],
            ['Uzay gemisi', 'X', 'Y'],
            ['İğne', '', '', 'adsız'],
            ['Terminal', '', 'Rolling fırdöndü', '', 'No:4', 'belki'],
            ['Yem', '', 'Karides', '', '', 'var'],
            ['İğne', 'Owner', 'SSW 5111', 'siyah nikel', '1/0', 'Var'],
        ]
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                ws.cell(i + 2, j + 1, v)
        fx = SHOTS + '/dolu.xlsx'; wb.save(fx)
        sh = await import_file(pg, fx)
        print('   özet:', sh[:400].replace('\n', ' | '))
        check('7 kalem eklenecek' in sh and '3 satır okunamadı' in sh and '1 kalem zaten envanterde' in sh, 'AJ: özet: 7 kalem, 3 okunamayan, 1 tekrar')
        check('Satır 9: “Uzay gemisi” kategorisi tanınmadı' in sh and 'Satır 10: ad yok' in sh and 'Satır 11: “belki” durumu tanınmadı' in sh, 'AJ: okunamayan satırlar numarasıyla')
        await overflow(pg, 'Excel özeti 390')
        await pg.set_viewport_size({'width': 320, 'height': 860}); await pg.wait_for_timeout(200); await overflow(pg, 'Excel özeti 320'); await pg.set_viewport_size({'width': 390, 'height': 860})
        await tall(pg, 1700, SHOTS + '/aj_ozet.png', '#sheet')
        await pg.locator('#sheet button[data-xl="ok"]').last.click(); await pg.wait_for_timeout(600)
        st = await pg.evaluate('JSON.parse(localStorage.getItem("awf-v1"))')
        its = {it['name']: it for it in st['inv']['items'].values()}
        print('   kalemler:', sorted(its))
        h = its.get('Owner SSW 5111') or {}
        check(h.get('cat') == 'hook' and h.get('sizes') == ['1/0', '2/0', '3/0'] and h.get('pattern') == 'Octopus / beak' and h.get('note') == 'mercan için' and h.get('stock', {}).get('2/0') == 'var', 'AJ: iğne: boylar, tip, not, durum')
        g = its.get('Gamakatsu Chinu') or {}
        check(g.get('sizes') == ['No:2', 'No:4'] and g.get('pattern') == 'Chinu' and g.get('stock', {}).get('No:4') == 'alinacak', 'AJ: küçük harf kategori ve tip, noktalı virgül, alınacak')
        f = its.get('Seaguar Grand Max FX') or {}
        print('   fc:', f)
        check(f.get('cat') == 'fc' and f.get('sizes') == [] and f.get('spec') == '0,33 mm · 0,40 mm' and f.get('stock', {}).get('-') == 'var', 'AJ: ipte boy özelliğe yazılır, boş durum Var')
        r = its.get('Testrod Slow 198') or {}
        check(r.get('rod') == {'type': 'slowjig', 'len': 1.98, 'max': 200, 'pe': 2}, 'AJ: kamış tipi, boy, en fazla, PE (%s)' % r.get('rod'))
        m = its.get('Testreel Lever 30') or {}
        check(m.get('reel') == {'type': 'cikrik', 'size': 30} and m.get('stock', {}).get('-') == 'alinacak', 'AJ: makine tipi ve boyu')
        k = its.get('Armut kurşun') or {}
        check(k.get('sizes') == ['60 g', '80 g'] and k.get('stock', {}).get('80 g') == 'yok', 'AJ: kurşun yok')
        y = its.get('Karides') or {}
        check(y.get('cat') == 'bait' and any(v == 'var' for kk, v in st['inv']['status'].items() if kk.startswith(y.get('id', '?') + '|')), 'AJ: yem var işareti')
        check(all(it.get('src') == 'kullanici' and 'xlRow' not in it for it in its.values()), 'AJ: kalemler kullanıcının envanterinde')
        await pg.click('.tabbar button[data-v="envanter"]'); await pg.wait_for_timeout(250)
        await pg.click('#invSeg button:has-text("İğneler")'); await pg.wait_for_timeout(250)
        check('SSW 5111' in await pg.inner_text('#invList'), 'AJ: iğneler listesinde görünür')
        await pg.click('.tabbar button[data-v="daha"]'); await pg.wait_for_timeout(200)
        await pg.click('#moreMenu button[data-open="alisveris"]'); await pg.wait_for_timeout(400)
        ls = await pg.inner_text('#sheet')
        print('   liste:', ls[:600].replace('\n', ' | '))
        check('Chinu' in ls and 'Lever 30' in ls, 'AJ: alınacaklar alışveriş listesinde')
        await pg.locator('#sheet [data-close]').first.click(); await pg.wait_for_timeout(200)
        # aynı dosya ikinci kez: hepsi tekrar
        sh = await import_file(pg, fx)
        print('   ikinci:', sh[:300].replace('\n', ' | '))
        check('0 kalem eklenecek' in sh and '8 kalem zaten envanterde' in sh, 'AJ: aynı dosya ikinci kez eklenmez')
        await pg.locator('#sheet [data-close]').first.click(); await pg.wait_for_timeout(200)
        # CSV (noktalı virgül, tırnak, BOM)
        csvp = SHOTS + '/envanter.csv'
        open(csvp, 'w', encoding='utf-8-sig').write('Kategori;Marka;Model / ad;Özellik;Boylar;Durum;Not\nJig / zoka;Major Craft;Jigpara Slow;pembe;"80 g, 100 g";Var;"tekne; derin"\nSahte yem;Rapala;X-Rap;;;Alınacak;\n')
        sh = await import_file(pg, csvp)
        check('2 kalem eklenecek' in sh, 'AJ: CSV okundu')
        await pg.locator('#sheet button[data-xl="ok"]').last.click(); await pg.wait_for_timeout(500)
        st = await pg.evaluate('JSON.parse(localStorage.getItem("awf-v1"))')
        jg = [it for it in st['inv']['items'].values() if it['name'] == 'Major Craft Jigpara Slow']
        check(jg and jg[0]['sizes'] == ['80 g', '100 g'] and jg[0]['note'] == 'tekne; derin' and jg[0]['cat'] == 'jig', 'AJ: CSV: tırnaklı hücreler')
        # bozuk dosya
        bad = SHOTS + '/bozuk.xlsx'; open(bad, 'wb').write(b'PK\x03\x04 bozuk')
        sh = await import_file(pg, bad)
        check('Dosya okunamadı' in sh, 'AJ: bozuk dosya uyarısı')
        await pg.locator('#sheet [data-close]').first.click(); await pg.wait_for_timeout(200)
        # başlıksız dosya
        nh = SHOTS + '/basliksiz.csv'; open(nh, 'w').write('a,b,c\n1,2,3\n')
        sh = await import_file(pg, nh)
        check('Başlık satırı bulunamadı' in sh, 'AJ: başlıksız dosya uyarısı')
        await pg.locator('#sheet [data-close]').first.click(); await pg.wait_for_timeout(200)
        await pg.click('.tabbar button[data-v="envanter"]'); await pg.wait_for_timeout(200)
        await pg.screenshot(path=SHOTS + '/aj_envanter.png')
        print('AJ genel errors:', errs); check(not errs, 'AJ: genel: hata yok')
        await ctx.close()
        # ---------------- özel yayın: paket kapısı ve eski adlar korunur
        ctx, pg, errs = await new_ctx(b)
        await pg.goto(BASE + '/dist/'); await pg.wait_for_timeout(1000)
        check(await pg.title() == 'Alyssa WF' and await pg.locator('button[data-wel="paket"]').count() == 0, 'AJ: özel: yeni kullanıcı paket düğmesi görmez')
        await pg.click('button[data-wel="bos"]'); await pg.wait_for_timeout(500)
        await pg.click('.tabbar button[data-v="rigler"]'); await pg.wait_for_timeout(300)
        check(await pg.locator('#v-rigler [data-open="paket"]').count() == 0, 'AJ: özel: yeni kullanıcı Rigler’de paket düğmesi görmez')
        ks = await lskeys(pg); check('alyssa-v1' in ks, 'AJ: özel: depolama alyssa- önekiyle (eski profil korunur)')
        await ctx.close()
        ctx, pg, errs = await new_ctx(b)
        await pg.goto(BASE + '/dist/?paket'); await pg.wait_for_timeout(1000)
        await pg.click('button[data-wel="paket"]'); await pg.fill('#wel-pass', KEY['pass_']); await pg.click('#welBtn'); await pg.wait_for_timeout(5000)
        await pg.click('.tabbar button[data-v="daha"]'); await pg.wait_for_timeout(300)
        ab = await pg.inner_text('#about')
        check(ab.startswith('Alyssa WF v0.29') and 'özel paket' in ab and await pg.locator('#moreMenu button[data-open="paket"]').count() == 1, 'AJ: özel: paket açılır, menüde görünür')
        dbs = await pg.evaluate('indexedDB.databases ? indexedDB.databases().then(d => d.map(x => x.name)) : []')
        check('kenny-av' in dbs, 'AJ: özel: av kayıtları eski veritabanında (%s)' % dbs)
        async with pg.expect_download() as dl_info:
            await pg.click('.tabbar button[data-v="envanter"]'); await pg.wait_for_timeout(200); await pg.click('#v-envanter button[data-xl="tpl"]')
        dl = await dl_info.value
        check(dl.suggested_filename == 'ornek-envanter.xlsx', 'AJ: özel: örnek dosya indirilir')
        print('AJ özel errors:', errs); check(not errs, 'AJ: özel: hata yok')
        await ctx.close(); await b.close()
    srv.terminate()
    print('SONUÇ:', 'HEPSİ GEÇTİ' if not FAIL else '%d HATA' % len(FAIL))
asyncio.run(main())
